from decimal import Decimal
from io import StringIO
from pathlib import Path
from unittest import mock

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import connection
from openpyxl import load_workbook
from PIL import Image

from apps.catalog.models import (
    Brand,
    Category,
    Product,
    ProductImage,
    ProductVariant,
    Shade,
    Size,
    SkinType,
)
from apps.catalog.services import stock
from apps.catalog.tests.factories import BrandFactory
from apps.catalog.workbook import importer
from apps.catalog.workbook.export import CHARCOAL, LIGHT_GREY, build_workbook
from apps.catalog.workbook.layout import (
    BRANDS,
    CATEGORIES,
    DATA_SHEETS,
    PRODUCTS,
    SHADES,
    SIZES,
    SKIN_TYPES,
    VARIANTS,
    CellValue,
    Sheet,
)

type Rows = dict[Sheet, list[list[CellValue]]]


def catalogue() -> Rows:
    return {
        BRANDS: [["Glow Lab", "Gentle skincare.", None, "Yes", 1]],
        CATEGORIES: [["Skincare", None, 1], ["Serums", "Skincare", 1]],
        SHADES: [["Warm Honey", "#D8A47F", 1]],
        SIZES: [["30 ml", 1], ["50 ml", 2]],
        SKIN_TYPES: [["Dry", 1], ["Oily", 2]],
        PRODUCTS: [
            [
                "Rose Serum",
                "Glow Lab",
                "Serums",
                2450,
                "Hydrates and calms.",
                "Dry, Oily",
                "Light",
                "Water, Glycerin",
                "Yes",
                None,
                None,
                1,
            ]
        ],
        VARIANTS: [
            ["Rose Serum", "GL-ROSE-30", "30 ml", None, 12, None, 2900],
            ["Rose Serum", "GL-ROSE-50", "50 ml", "Warm Honey", 5, 3100.5, None],
        ],
    }


def write(tmp_path: Path, rows: Rows, name: str = "catalogue.xlsx") -> Path:
    path = tmp_path / name
    build_workbook(
        {sheet: [tuple(row) for row in sheet_rows] for sheet, sheet_rows in rows.items()}
    ).save(path)
    return path


def run_import(path: Path, *args: str) -> str:
    out = StringIO()
    call_command("import_catalogue", str(path), *args, stdout=out, stderr=StringIO())
    return out.getvalue()


def failed_import(path: Path, *args: str) -> str:
    err = StringIO()
    with pytest.raises(CommandError):
        call_command("import_catalogue", str(path), *args, stdout=StringIO(), stderr=err)
    return err.getvalue()


def png(path: Path, image_format: str = "PNG") -> Path:
    Image.new("RGB", (8, 8), "#D8A47F").save(path, format=image_format)
    return path


@pytest.fixture
def media(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path / "media"


@pytest.fixture
def images(tmp_path):
    folder = tmp_path / "images"
    folder.mkdir()
    for name in ("front.png", "back.png", "new.png", "logo.png"):
        png(folder / name)
    return folder


# --- The template -------------------------------------------------------------


@pytest.fixture
def template(tmp_path):
    path = tmp_path / "template.xlsx"
    call_command("export_catalogue", str(path), "--template", stdout=StringIO())
    return load_workbook(path)


def test_template_has_the_read_me_and_data_sheets_in_order(template):
    assert template.sheetnames == ["Read me", *(sheet.title for sheet in DATA_SHEETS)]


def test_read_me_explains_the_legend_rules_and_contact(template):
    text = "\n".join(str(cell.value) for cell in template["Read me"]["A"] if cell.value)

    assert "Row 2 on every sheet is an example" in text
    assert "#D8A47F" in text
    assert "first image" in text
    assert "Contact:" in text


@pytest.mark.parametrize("sheet", DATA_SHEETS, ids=lambda sheet: sheet.title)
def test_every_sheet_has_styled_commented_headers_and_frozen_row(template, sheet):
    worksheet = template[sheet.title]

    assert [cell.value for cell in worksheet[1]] == [column.header for column in sheet.columns]
    for cell, column in zip(worksheet[1], sheet.columns, strict=True):
        assert cell.font.name == "Arial"
        assert cell.font.bold
        assert cell.fill.start_color.rgb.endswith(CHARCOAL if column.required else LIGHT_GREY)
        assert cell.comment.text == column.comment
    assert worksheet.freeze_panes == "A2"


@pytest.mark.parametrize("sheet", DATA_SHEETS, ids=lambda sheet: sheet.title)
def test_row_two_is_a_grey_italic_example(template, sheet):
    worksheet = template[sheet.title]

    assert worksheet["A2"].value == sheet.example[0]
    assert worksheet["A2"].font.italic
    assert "ignored" in worksheet["A2"].comment.text


def test_template_prefills_the_skin_types(template):
    names = [row[0] for row in template["Skin types"].iter_rows(min_row=3, values_only=True)]

    assert names == ["Normal", "Dry", "Oily", "Combination", "Sensitive", "Mature"]


def _validations(worksheet):
    return {
        str(validation.sqref): validation
        for validation in worksheet.data_validations.dataValidation
    }


def test_references_are_dropdowns_of_the_other_sheets(template):
    products = _validations(template["Products"])
    variants = _validations(template["Variants"])
    categories = _validations(template["Categories"])

    assert products["B3:B1000"].formula1 == "'Brands'!$A$3:$A$1000"
    assert products["C3:C1000"].formula1 == "'Categories'!$A$3:$A$1000"
    assert variants["C3:C1000"].formula1 == "'Sizes'!$A$3:$A$1000"
    assert variants["D3:D1000"].formula1 == "'Shades'!$A$3:$A$1000"
    assert categories["B3:B1000"].formula1 == "'Categories'!$A$3:$A$1000"
    assert products["B3:B1000"].showErrorMessage
    assert "Brands sheet" in products["B3:B1000"].error


def test_yes_no_numbers_and_prices_are_validated(template):
    brands = _validations(template["Brands"])
    products = _validations(template["Products"])
    variants = _validations(template["Variants"])

    assert brands["D3:D1000"].formula1 == '"Yes,No"'
    assert products["I3:I1000"].formula1 == '"Yes,No"'
    assert (variants["E3:E1000"].type, variants["E3:E1000"].formula1) == ("whole", "0")
    assert brands["E3:E1000"].type == "whole"
    assert (products["D3:D1000"].type, products["D3:D1000"].operator) == (
        "decimal",
        "greaterThanOrEqual",
    )
    assert variants["F3:F1000"].type == "decimal"


def test_the_workbook_has_no_formulas(template):
    for worksheet in template.worksheets:
        for row in worksheet.iter_rows(values_only=True):
            assert not any(isinstance(value, str) and value.startswith("=") for value in row)


# --- Import -------------------------------------------------------------------


@pytest.mark.django_db
def test_import_creates_the_catalogue(tmp_path, media):
    output = run_import(write(tmp_path, catalogue()))

    product = Product.objects.get(slug="rose-serum")
    assert product.brand.name == "Glow Lab"
    assert product.category.parent == Category.objects.get(name="Skincare")
    assert product.base_price == Decimal("2450.00")
    assert product.is_published
    assert {s.name for s in product.skin_types.all()} == {"Dry", "Oily"}
    large = ProductVariant.objects.get(sku="GL-ROSE-50")
    assert (large.size.name, large.shade, large.stock_quantity) == (
        "50 ml",
        Shade.objects.get(name="Warm Honey"),
        5,
    )
    assert large.price_override == Decimal("3100.50")
    assert ProductVariant.objects.get(sku="GL-ROSE-30").shade is None
    assert "Products: 1 created, 0 updated, 0 unchanged" in output
    assert "Variants: 2 created, 0 updated, 0 unchanged" in output


@pytest.mark.django_db
def test_the_example_row_is_ignored(tmp_path, media):
    path = tmp_path / "template.xlsx"
    call_command("export_catalogue", str(path), "--template", stdout=StringIO())

    run_import(path)

    assert not Brand.objects.exists()
    assert not Product.objects.exists()
    assert list(SkinType.objects.values_list("name", flat=True)) == [
        "Normal",
        "Dry",
        "Oily",
        "Combination",
        "Sensitive",
        "Mature",
    ]


@pytest.mark.django_db
def test_export_then_import_changes_nothing(tmp_path, media, images):
    rows = catalogue()
    rows[BRANDS][0][2] = "logo.png"
    rows[PRODUCTS][0][9] = "front.png, back.png"
    run_import(write(tmp_path, rows), "--images", str(images))
    before = _snapshot()

    exported = tmp_path / "export.xlsx"
    call_command("export_catalogue", str(exported), stdout=StringIO())
    first = run_import(exported)
    second = run_import(exported)

    assert _snapshot() == before
    for output in (first, second):
        for sheet in DATA_SHEETS:
            assert f"{sheet.title}: 0 created, 0 updated," in output
        assert "Images: 0 uploaded" in output


def _snapshot():
    return {
        "brands": list(Brand.objects.values_list("name", "description", "logo", "is_active")),
        "categories": list(Category.objects.values_list("name", "parent__name", "sort_order")),
        "shades": list(Shade.objects.values_list("name", "hex_code")),
        "sizes": list(Size.objects.values_list("name", "sort_order")),
        "products": list(
            Product.objects.values_list(
                "name", "slug", "base_price", "is_published", "description", "updated_at"
            )
        ),
        "variants": list(
            ProductVariant.objects.order_by("sku").values_list(
                "sku", "stock_quantity", "price_override", "compare_at_price", "updated_at"
            )
        ),
        "images": list(ProductImage.objects.values_list("image", "is_primary")),
    }


@pytest.mark.django_db
def test_second_import_updates_changed_fields_only(tmp_path, media):
    run_import(write(tmp_path, catalogue()))
    rows = catalogue()
    rows[BRANDS][0][1] = "Now with SPF."
    rows[PRODUCTS][0][3] = "2600"
    rows[PRODUCTS][0][5] = "Dry"
    rows[VARIANTS][1][5] = None

    output = run_import(write(tmp_path, rows, "changed.xlsx"))

    assert Brand.objects.get().description == "Now with SPF."
    product = Product.objects.get()
    assert product.base_price == Decimal("2600.00")
    assert [s.name for s in product.skin_types.all()] == ["Dry"]
    assert ProductVariant.objects.get(sku="GL-ROSE-50").price_override is None
    assert "Brands: 0 created, 1 updated, 0 unchanged" in output
    assert "Products: 0 created, 1 updated, 0 unchanged" in output
    assert "Variants: 0 created, 1 updated, 1 unchanged" in output
    assert "Shades: 0 created, 0 updated, 1 unchanged" in output


@pytest.mark.django_db
def test_stock_is_set_through_set_variant_stock(tmp_path, media):
    run_import(write(tmp_path, catalogue()))
    rows = catalogue()
    rows[VARIANTS][0][4] = 40

    with (
        mock.patch.object(stock, "set_variant_stock", wraps=stock.set_variant_stock) as spy,
        mock.patch("apps.catalog.services.products.set_variant_stock", spy),
    ):
        run_import(write(tmp_path, rows, "restocked.xlsx"))

    variant = ProductVariant.objects.get(sku="GL-ROSE-30")
    spy.assert_called_once_with(variant=variant, quantity=40)
    assert variant.stock_quantity == 40


@pytest.mark.django_db
def test_new_variants_get_their_stock_through_set_variant_stock(tmp_path, media):
    with mock.patch.object(importer, "set_variant_stock", wraps=stock.set_variant_stock) as spy:
        run_import(write(tmp_path, catalogue()))

    assert sorted(call.kwargs["quantity"] for call in spy.call_args_list) == [5, 12]


@pytest.mark.django_db
def test_rows_missing_from_the_workbook_are_reported_and_kept(tmp_path, media):
    BrandFactory(name="Old Brand")

    output = run_import(write(tmp_path, catalogue()))

    assert Brand.objects.filter(name="Old Brand").exists()
    assert "1 not in workbook, left unchanged: Old Brand" in output


@pytest.mark.django_db
def test_dry_run_reports_and_writes_nothing(tmp_path, media):
    output = run_import(write(tmp_path, catalogue()), "--dry-run")

    assert "Dry run" in output
    assert "Products: 1 created" in output
    assert not Brand.objects.exists()
    assert not Product.objects.exists()


@pytest.mark.django_db
def test_import_accepts_loose_spellings_and_numbers_as_text(tmp_path, media):
    rows = catalogue()
    rows[BRANDS][0] = ["  Glow Lab  ", None, None, "n", "3"]
    rows[PRODUCTS][0][3] = "2450.50"
    rows[PRODUCTS][0][8] = "YES"
    rows[VARIANTS][0][4] = "12"

    run_import(write(tmp_path, rows))

    brand = Brand.objects.get()
    assert (brand.name, brand.is_active, brand.sort_order) == ("Glow Lab", False, 3)
    assert Product.objects.get().base_price == Decimal("2450.50")
    assert ProductVariant.objects.get(sku="GL-ROSE-30").stock_quantity == 12


@pytest.mark.django_db
def test_import_matches_existing_rows_already_in_the_database(tmp_path, media):
    run_import(write(tmp_path, catalogue()))
    rows: Rows = {sheet: [] for sheet in DATA_SHEETS}
    rows[PRODUCTS] = [catalogue()[PRODUCTS][0]]
    rows[VARIANTS] = [["Rose Serum", "GL-ROSE-100", "50 ml", None, 3, None]]

    output = run_import(write(tmp_path, rows, "more.xlsx"))

    assert ProductVariant.objects.filter(product__slug="rose-serum").count() == 3
    assert "Variants: 1 created, 0 updated, 0 unchanged" in output


# --- Validation ---------------------------------------------------------------


def _set(rows: Rows, sheet: Sheet, index: int, column: str, value: CellValue) -> None:
    headers = [column.header for column in sheet.columns]
    rows[sheet][index][headers.index(column)] = value


def _append(rows: Rows, sheet: Sheet, row: list[CellValue]) -> None:
    rows[sheet].append(row)


INVALID = {
    "required": (
        lambda rows: _set(rows, PRODUCTS, 0, "Brand", None),
        'Sheet "Products", row 3, column "Brand": is required',
    ),
    "unknown brand": (
        lambda rows: _set(rows, PRODUCTS, 0, "Brand", "Lumiere"),
        'Sheet "Products", row 3, column "Brand": "Lumiere" is not on the Brands sheet',
    ),
    "unknown category": (
        lambda rows: _set(rows, PRODUCTS, 0, "Category", "Toners"),
        'Sheet "Products", row 3, column "Category": "Toners" is not on the Categories sheet',
    ),
    "unknown skin type": (
        lambda rows: _set(rows, PRODUCTS, 0, "Skin types", "Dry, Shiny"),
        'Sheet "Products", row 3, column "Skin types": "Shiny" is not on the Skin types sheet',
    ),
    "unknown size": (
        lambda rows: _set(rows, VARIANTS, 0, "Size", "75 ml"),
        'Sheet "Variants", row 3, column "Size": "75 ml" is not on the Sizes sheet',
    ),
    "unknown shade": (
        lambda rows: _set(rows, VARIANTS, 0, "Shade", "Teal"),
        'Sheet "Variants", row 3, column "Shade": "Teal" is not on the Shades sheet',
    ),
    "duplicate brand": (
        lambda rows: _append(rows, BRANDS, ["Glow Lab", None, None, None, None]),
        'Sheet "Brands", row 4, column "Name": "Glow Lab" is already on row 3',
    ),
    "duplicate shade": (
        lambda rows: _append(rows, SHADES, ["Warm Honey", "#000000", None]),
        'Sheet "Shades", row 4, column "Name": "Warm Honey" is already on row 3',
    ),
    "duplicate size": (
        lambda rows: _append(rows, SIZES, ["30 ml", None]),
        'Sheet "Sizes", row 5, column "Name": "30 ml" is already on row 3',
    ),
    "duplicate skin type": (
        lambda rows: _append(rows, SKIN_TYPES, ["Dry", None]),
        'Sheet "Skin types", row 5, column "Name": "Dry" is already on row 3',
    ),
    "duplicate product": (
        lambda rows: _append(rows, PRODUCTS, list(rows[PRODUCTS][0])),
        'Sheet "Products", row 4, column "Product name": "Rose Serum" is already on row 3',
    ),
    "duplicate sku": (
        lambda rows: _append(rows, VARIANTS, ["Rose Serum", "GL-ROSE-30", "50 ml", None, 1, None]),
        'Sheet "Variants", row 5, column "SKU": "GL-ROSE-30" is already on row 3',
    ),
    "duplicate size and shade": (
        lambda rows: _append(rows, VARIANTS, ["Rose Serum", "GL-ROSE-30B", "30 ml", None, 1, None]),
        'Sheet "Variants", row 5, column "Size": "Rose Serum" already has a variant in 30 ml, '
        "no shade on row 3",
    ),
    "bad hex": (
        lambda rows: _set(rows, SHADES, 0, "Hex colour", "D8A47F"),
        'Sheet "Shades", row 3, column "Hex colour": "D8A47F" is not a hex colour like #D8A47F',
    ),
    "price with rupees": (
        lambda rows: _set(rows, PRODUCTS, 0, "Base price (NPR)", "Rs 2,450"),
        'Sheet "Products", row 3, column "Base price (NPR)": "Rs 2,450" is not a price',
    ),
    "negative price": (
        lambda rows: _set(rows, PRODUCTS, 0, "Base price (NPR)", -1),
        'column "Base price (NPR)": -1 is negative',
    ),
    "three decimals": (
        lambda rows: _set(rows, VARIANTS, 1, "Price override (NPR)", 3100.555),
        'Sheet "Variants", row 4, column "Price override (NPR)": 3100.555 has more than 2 '
        "decimal places",
    ),
    "zero override": (
        lambda rows: _set(rows, VARIANTS, 1, "Price override (NPR)", 0),
        'column "Price override (NPR)": must be more than 0',
    ),
    "compare-at not above the base price": (
        lambda rows: _set(rows, VARIANTS, 0, "Compare-at price (NPR)", 2450),
        'Sheet "Variants", row 3, column "Compare-at price (NPR)": must be more than the '
        "price, 2450.00",
    ),
    "compare-at not above the override": (
        lambda rows: _set(rows, VARIANTS, 1, "Compare-at price (NPR)", 3000),
        'Sheet "Variants", row 4, column "Compare-at price (NPR)": must be more than the '
        "price, 3100.50",
    ),
    "fractional stock": (
        lambda rows: _set(rows, VARIANTS, 0, "Stock", 12.5),
        'Sheet "Variants", row 3, column "Stock": "12.5" is not a whole number',
    ),
    "negative stock": (
        lambda rows: _set(rows, VARIANTS, 0, "Stock", -2),
        'column "Stock": -2 is negative',
    ),
    "yes or no": (
        lambda rows: _set(rows, PRODUCTS, 0, "Published", "maybe"),
        'Sheet "Products", row 3, column "Published": "maybe" is not Yes or No',
    ),
    "two levels deep": (
        lambda rows: _append(rows, CATEGORIES, ["Night serums", "Serums", None]),
        'Sheet "Categories", row 5, column "Parent category": "Serums" is itself a '
        "sub-category; categories are one level deep only",
    ),
    "unknown parent": (
        lambda rows: _set(rows, CATEGORIES, 1, "Parent category", "Skin"),
        'Sheet "Categories", row 4, column "Parent category": "Skin" is not on the '
        "Categories sheet",
    ),
    "variant of unknown product": (
        lambda rows: _set(rows, VARIANTS, 0, "Product name", "Rose Serun"),
        'Sheet "Variants", row 3, column "Product name": "Rose Serun" is not on the Products sheet',
    ),
    "published without variants": (
        lambda rows: rows[VARIANTS].clear(),
        'Sheet "Products", row 3, column "Published": A product needs at least one variant',
    ),
    "missing image": (
        lambda rows: _set(rows, PRODUCTS, 0, "Image files", "front.png, side.png"),
        'Sheet "Products", row 3, column "Image files": "side.png" is not in the images folder',
    ),
    "image path": (
        lambda rows: _set(rows, BRANDS, 0, "Logo file", "../logo.png"),
        'Sheet "Brands", row 3, column "Logo file": "../logo.png" must be a file name',
    ),
    "not an image": (
        lambda rows: _set(rows, PRODUCTS, 0, "Image files", "notes.png"),
        '"notes.png" is not a JPEG, PNG or WebP image',
    ),
    "wrong image type": (
        lambda rows: _set(rows, PRODUCTS, 0, "Image files", "anim.gif"),
        '"anim.gif" is not a JPEG, PNG or WebP image',
    ),
    "image too large": (
        lambda rows: _set(rows, PRODUCTS, 0, "Image files", "huge.png"),
        '"huge.png" is 5.0 MB; the limit is 5 MB',
    ),
    "name too long": (
        lambda rows: _set(rows, SIZES, 0, "Name", "x" * 51),
        'Sheet "Sizes", row 3, column "Name": is 51 characters long; the limit is 50',
    ),
}


@pytest.mark.django_db
@pytest.mark.parametrize("case", INVALID.values(), ids=INVALID.keys())
def test_each_problem_is_reported_by_sheet_row_and_column_and_nothing_is_written(
    tmp_path, media, images, case
):
    mutate, message = case
    (images / "notes.png").write_text("not an image")
    png(images / "anim.gif", "GIF")
    (images / "huge.png").write_bytes(b"\0" * (5 * 1024 * 1024 + 1))
    rows = catalogue()
    mutate(rows)

    errors = failed_import(write(tmp_path, rows), "--images", str(images))

    assert message in errors
    assert not Brand.objects.exists()
    assert not Category.objects.exists()
    assert not Product.objects.exists()
    assert not SkinType.objects.exists()


@pytest.mark.django_db
def test_every_problem_is_reported_at_once(tmp_path, media):
    rows = catalogue()
    _set(rows, PRODUCTS, 0, "Brand", "Lumiere")
    _set(rows, SHADES, 0, "Hex colour", "red")
    _set(rows, VARIANTS, 0, "Stock", -1)

    errors = failed_import(write(tmp_path, rows))

    assert len(errors.strip().splitlines()) == 3


@pytest.mark.django_db
def test_an_image_listed_without_an_images_folder_is_a_problem(tmp_path, media):
    rows = catalogue()
    _set(rows, PRODUCTS, 0, "Image files", "front.png")

    errors = failed_import(write(tmp_path, rows))

    assert '"front.png" needs the images folder: pass --images <folder>' in errors


@pytest.mark.django_db
def test_a_sku_of_another_product_is_a_problem(tmp_path, media):
    run_import(write(tmp_path, catalogue()))
    rows = catalogue()
    rows[PRODUCTS].append(["Night Cream", "Glow Lab", "Serums", 1800, *[None] * 8])
    _set(rows, VARIANTS, 0, "Product name", "Night Cream")

    errors = failed_import(write(tmp_path, rows, "moved.xlsx"))

    assert '"GL-ROSE-30" already belongs to "Rose Serum"' in errors


@pytest.mark.django_db
def test_a_missing_sheet_or_column_is_a_problem(tmp_path, media):
    path = write(tmp_path, catalogue())
    workbook = load_workbook(path)
    del workbook["Sizes"]
    workbook["Brands"]["A1"] = "Brand name"
    workbook.save(path)

    errors = failed_import(path)

    assert 'Sheet "Sizes": the sheet is missing from the workbook' in errors
    assert 'Sheet "Brands": the column "Name" is missing from row 1' in errors


@pytest.mark.django_db
def test_a_workbook_from_before_the_compare_at_column_still_imports(tmp_path, media):
    run_import(write(tmp_path, catalogue()))
    exported = tmp_path / "export.xlsx"
    call_command("export_catalogue", str(exported), stdout=StringIO())
    workbook = load_workbook(exported)
    variants = workbook[VARIANTS.title]
    headers = [cell.value for cell in variants[1]]
    variants.delete_cols(headers.index("Compare-at price (NPR)") + 1)
    workbook.save(exported)
    ProductVariant.objects.update(compare_at_price=None)

    output = run_import(exported)

    for sheet in DATA_SHEETS:
        assert f"{sheet.title}: 0 created, 0 updated," in output
    assert not ProductVariant.objects.filter(compare_at_price__isnull=False).exists()


@pytest.mark.django_db
def test_a_missing_optional_column_reads_as_blank_and_an_unknown_one_is_ignored(tmp_path, media):
    path = write(tmp_path, catalogue())
    workbook = load_workbook(path)
    workbook["Brands"]["D1"] = "Is active"
    workbook.save(path)

    run_import(path)

    assert Brand.objects.get().is_active is True


def test_a_file_that_is_not_a_workbook_is_refused(tmp_path):
    path = tmp_path / "catalogue.xlsx"
    path.write_text("name,price")

    with pytest.raises(CommandError, match=r"not an \.xlsx workbook"):
        call_command("import_catalogue", str(path), stdout=StringIO())


# --- Images -------------------------------------------------------------------


@pytest.mark.django_db
def test_images_and_logos_are_uploaded_first_image_primary(tmp_path, media, images):
    rows = catalogue()
    _set(rows, BRANDS, 0, "Logo file", "logo.png")
    _set(rows, PRODUCTS, 0, "Image files", "front.png, back.png")
    _set(rows, PRODUCTS, 0, "Image alt text", "Rose Serum bottle")

    output = run_import(write(tmp_path, rows), "--images", str(images))

    assert str(Brand.objects.get().logo.name).startswith("brands/logo")
    product_images = list(Product.objects.get().images.all())
    assert [Path(str(image.image.name)).name for image in product_images] == [
        "front.png",
        "back.png",
    ]
    assert [image.is_primary for image in product_images] == [True, False]
    assert {image.alt_text for image in product_images} == {"Rose Serum bottle"}
    assert "Images: 3 uploaded" in output


@pytest.mark.django_db
def test_existing_images_are_kept_unless_replacing(tmp_path, media, images):
    rows = catalogue()
    _set(rows, PRODUCTS, 0, "Image files", "front.png, back.png")
    run_import(write(tmp_path, rows), "--images", str(images))
    _set(rows, PRODUCTS, 0, "Image files", "new.png")
    path = write(tmp_path, rows, "new.xlsx")

    kept = run_import(path, "--images", str(images))
    assert Product.objects.get().images.count() == 2
    assert "Images: 0 uploaded" in kept

    replaced = run_import(path, "--images", str(images), "--replace-images")
    product_images = list(Product.objects.get().images.all())
    assert [Path(str(image.image.name)).name for image in product_images] == ["new.png"]
    assert product_images[0].is_primary
    assert "Images: 1 uploaded" in replaced


@pytest.mark.django_db(transaction=True)
def test_images_upload_after_the_catalogue_is_committed(tmp_path, media, images):
    rows = catalogue()
    _set(rows, PRODUCTS, 0, "Image files", "front.png")
    seen = []

    def add_product_image(**kwargs):
        seen.append(
            (connection.in_atomic_block, Product.objects.filter(slug="rose-serum").exists())
        )
        return original(**kwargs)

    original = importer.add_product_image
    with mock.patch.object(importer, "add_product_image", add_product_image):
        run_import(write(tmp_path, rows), "--images", str(images))

    assert seen == [(False, True)]


@pytest.mark.django_db
def test_a_failed_upload_is_reported_without_stopping_the_rest(tmp_path, media, images):
    rows = catalogue()
    _set(rows, PRODUCTS, 0, "Image files", "front.png, back.png")
    original = importer.add_product_image

    def flaky(**kwargs):
        if kwargs["image"].name == "front.png":
            raise OSError("connection reset")
        return original(**kwargs)

    err = StringIO()
    with (
        mock.patch.object(importer, "add_product_image", flaky),
        pytest.raises(CommandError, match="1 upload"),
    ):
        call_command(
            "import_catalogue",
            str(write(tmp_path, rows)),
            "--images",
            str(images),
            stdout=StringIO(),
            stderr=err,
        )

    assert 'Product "Rose Serum", "front.png": connection reset' in err.getvalue()
    product_images = list(Product.objects.get().images.all())
    assert [Path(str(image.image.name)).name for image in product_images] == ["back.png"]
    assert product_images[0].is_primary


@pytest.mark.django_db
def test_import_sets_and_clears_the_compare_at_price(tmp_path, media):
    run_import(write(tmp_path, catalogue()))
    assert ProductVariant.objects.get(sku="GL-ROSE-30").compare_at_price == Decimal("2900.00")

    rows = catalogue()
    _set(rows, VARIANTS, 0, "Compare-at price (NPR)", None)
    output = run_import(write(tmp_path, rows, "cleared.xlsx"))

    assert ProductVariant.objects.get(sku="GL-ROSE-30").compare_at_price is None
    assert "Variants: 0 created, 1 updated, 1 unchanged" in output


@pytest.mark.django_db
def test_a_sale_can_start_in_the_same_import_that_lowers_the_base_price(tmp_path, media):
    rows = catalogue()
    _set(rows, VARIANTS, 0, "Compare-at price (NPR)", None)
    run_import(write(tmp_path, rows))

    rows = catalogue()
    _set(rows, PRODUCTS, 0, "Base price (NPR)", 2000)
    _set(rows, VARIANTS, 0, "Compare-at price (NPR)", 2450)
    run_import(write(tmp_path, rows, "sale.xlsx"))

    variant = ProductVariant.objects.get(sku="GL-ROSE-30")
    assert (variant.price, variant.compare_at_price) == (Decimal("2000.00"), Decimal("2450.00"))
    assert variant.on_sale is True


@pytest.mark.django_db
def test_the_compare_at_is_checked_against_a_product_already_in_the_database(tmp_path, media):
    run_import(write(tmp_path, catalogue()))
    rows = catalogue()
    rows[PRODUCTS].clear()
    _append(rows, VARIANTS, ["Rose Serum", "GL-ROSE-50B", "50 ml", None, 1, None, 2000])

    errors = failed_import(write(tmp_path, rows, "db.xlsx"))

    assert (
        'Sheet "Variants", row 5, column "Compare-at price (NPR)": must be more than the '
        "price, 2450.00" in errors
    )
