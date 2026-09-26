from collections.abc import Iterable, Mapping
from pathlib import PurePosixPath

from django.db.models import F
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from apps.catalog.models import (
    Brand,
    Category,
    Product,
    ProductVariant,
    Shade,
    Size,
    SkinType,
)
from apps.catalog.workbook.layout import (
    BRANDS,
    CATEGORIES,
    DATA_SHEETS,
    EXAMPLE_ROW,
    FIRST_DATA_ROW,
    FONT_NAME,
    LAST_VALIDATED_ROW,
    NO,
    PREFILLED_SKIN_TYPES,
    PRODUCTS,
    READ_ME,
    READ_ME_TITLE,
    SHADES,
    SIZES,
    SKIN_TYPES,
    VARIANTS,
    YES,
    CellValue,
    Column,
    Kind,
    Sheet,
)

type Rows = Mapping[Sheet, Iterable[tuple[CellValue, ...]]]

CHARCOAL = "2B2B2B"
LIGHT_GREY = "D9D9D9"
EXAMPLE_GREY = "808080"
COMMENT_AUTHOR = "TrueLux"


def template_rows() -> Rows:
    return {SKIN_TYPES: PREFILLED_SKIN_TYPES}


def catalogue_rows() -> Rows:
    return {
        BRANDS: [
            (
                brand.name,
                brand.description,
                _basename(brand.logo.name),
                _yes_no(brand.is_active),
                brand.sort_order,
            )
            for brand in Brand.objects.all()
        ],
        CATEGORIES: [
            (category.name, category.parent.name if category.parent else None, category.sort_order)
            for category in Category.objects.select_related("parent").order_by(
                F("parent__sort_order").asc(nulls_first=True), "parent__name", "sort_order", "name"
            )
        ],
        SHADES: [(shade.name, shade.hex_code, shade.sort_order) for shade in Shade.objects.all()],
        SIZES: [(size.name, size.sort_order) for size in Size.objects.all()],
        SKIN_TYPES: [(skin.name, skin.sort_order) for skin in SkinType.objects.all()],
        PRODUCTS: [_product_row(product) for product in _products()],
        VARIANTS: [
            (
                variant.product.name,
                variant.sku,
                variant.size.name,
                variant.shade.name if variant.shade else None,
                variant.stock_quantity,
                float(variant.price_override) if variant.price_override is not None else None,
            )
            for variant in ProductVariant.objects.select_related(
                "product__brand", "size", "shade"
            ).order_by(
                "product__brand__sort_order",
                "product__brand__name",
                "product__name",
                "size__sort_order",
                "shade__sort_order",
                "sku",
            )
        ],
    }


def _products() -> Iterable[Product]:
    return (
        Product.objects.select_related("brand", "category")
        .prefetch_related("skin_types", "images")
        .order_by("brand__sort_order", "brand__name", "name")
    )


def _product_row(product: Product) -> tuple[CellValue, ...]:
    # The workbook's first image is the main one, whatever its sort order.
    images = sorted(
        product.images.all(), key=lambda image: (not image.is_primary, image.sort_order)
    )
    return (
        product.name,
        product.brand.name,
        product.category.name,
        float(product.base_price),
        product.description,
        ", ".join(skin_type.name for skin_type in product.skin_types.all()),
        product.skin_feel,
        product.key_ingredients,
        _yes_no(product.is_published),
        ", ".join(_basename(image.image.name) for image in images),
        images[0].alt_text if images else None,
        product.sort_order,
    )


def _basename(name: str | None) -> str:
    return PurePosixPath(name).name if name else ""


def _yes_no(value: bool) -> str:
    return YES if value else NO


def build_workbook(rows: Rows) -> Workbook:
    workbook = Workbook()
    # openpyxl has no public setter for the default style, and cells the client types
    # into later take the workbook's Normal font rather than any cell's.
    workbook._named_styles["Normal"].font = Font(name=FONT_NAME, size=10)  # type: ignore[attr-defined]
    _write_read_me(workbook.create_sheet(READ_ME_TITLE))
    del workbook["Sheet"]
    for sheet in DATA_SHEETS:
        _write_sheet(workbook.create_sheet(sheet.title), sheet, rows.get(sheet, ()))
    return workbook


def _write_read_me(worksheet: Worksheet) -> None:
    worksheet.column_dimensions["A"].width = 110
    styles = {
        "title": (Font(name=FONT_NAME, size=16, bold=True), None),
        "heading": (Font(name=FONT_NAME, size=12, bold=True), None),
        "text": (Font(name=FONT_NAME, size=10), None),
        "required": (Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF"), CHARCOAL),
        "optional": (Font(name=FONT_NAME, size=10, bold=True), LIGHT_GREY),
        "example": (Font(name=FONT_NAME, size=10, italic=True, color=EXAMPLE_GREY), None),
    }
    for row, (style, text) in enumerate(READ_ME, start=1):
        cell = worksheet.cell(row=row, column=1, value=text)
        font, fill = styles[style]
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            cell.fill = PatternFill("solid", start_color=fill)


def _write_sheet(worksheet: Worksheet, sheet: Sheet, rows: Iterable[tuple[CellValue, ...]]) -> None:
    for index, column in enumerate(sheet.columns, start=1):
        letter = get_column_letter(index)
        worksheet.column_dimensions[letter].width = column.width
        _style_header(worksheet.cell(row=1, column=index, value=column.header), column)
        _add_validation(worksheet, letter, column)

    for index, value in enumerate(sheet.example, start=1):
        cell = worksheet.cell(row=EXAMPLE_ROW, column=index, value=value)
        cell.font = Font(name=FONT_NAME, size=10, italic=True, color=EXAMPLE_GREY)
    worksheet.cell(row=EXAMPLE_ROW, column=1).comment = Comment(
        "Example row: ignored when the workbook is loaded. Start your data on row 3.",
        COMMENT_AUTHOR,
    )

    for row_number, row in enumerate(rows, start=FIRST_DATA_ROW):
        for index, value in enumerate(row, start=1):
            worksheet.cell(row=row_number, column=index, value=value if value != "" else None)

    worksheet.freeze_panes = "A2"


def _style_header(cell: Cell, column: Column) -> None:
    if column.required:
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=CHARCOAL)
    else:
        cell.font = Font(name=FONT_NAME, size=10, bold=True)
        cell.fill = PatternFill("solid", start_color=LIGHT_GREY)
    cell.alignment = Alignment(vertical="center")
    cell.comment = Comment(column.comment, COMMENT_AUTHOR, width=320, height=110)


def _add_validation(worksheet: Worksheet, letter: str, column: Column) -> None:
    validation = _validation_for(column)
    if validation is None:
        return
    validation.showErrorMessage = True
    validation.errorTitle = f"Check {column.header}"
    worksheet.add_data_validation(validation)
    validation.add(f"{letter}{FIRST_DATA_ROW}:{letter}{LAST_VALIDATED_ROW}")


def _validation_for(column: Column) -> DataValidation | None:
    if column.choices_from:
        source = f"'{column.choices_from}'!$A${FIRST_DATA_ROW}:$A${LAST_VALIDATED_ROW}"
        return DataValidation(
            type="list",
            formula1=source,
            allow_blank=True,
            error=f"Pick a name from the {column.choices_from} sheet, or add it there first.",
        )
    if column.kind is Kind.YES_NO:
        return DataValidation(
            type="list", formula1=f'"{YES},{NO}"', allow_blank=True, error="Choose Yes or No."
        )
    if column.kind is Kind.WHOLE:
        return DataValidation(
            type="whole",
            operator="greaterThanOrEqual",
            formula1="0",
            allow_blank=True,
            error="Enter a whole number, 0 or more, e.g. 12.",
        )
    if column.kind is Kind.PRICE:
        return DataValidation(
            type="decimal",
            operator="greaterThanOrEqual",
            formula1="0",
            allow_blank=True,
            error='Enter the price in NPR as a plain number, e.g. 2450. No commas or "Rs".',
        )
    return None
