import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from apps.catalog.workbook.layout import (
    DATA_SHEETS,
    FIRST_DATA_ROW,
    HEADER_ROW,
    Column,
    Kind,
    Sheet,
)

HEX_CODE = re.compile(r"^#[0-9A-Fa-f]{6}$")
TRUE_SPELLINGS = frozenset({"yes", "y"})
FALSE_SPELLINGS = frozenset({"no", "n"})
# DecimalField(max_digits=10, decimal_places=2) holds at most eight digits before the point.
MAX_PRICE = Decimal("99999999.99")


@dataclass(frozen=True)
class Problem:
    sheet: str
    message: str
    row: int | None = None
    column: str | None = None

    def __str__(self) -> str:
        location = f'Sheet "{self.sheet}"'
        if self.row is not None:
            location += f", row {self.row}"
        if self.column is not None:
            location += f', column "{self.column}"'
        return f"{location}: {self.message}"


@dataclass
class Row:
    sheet: Sheet
    number: int
    values: dict[str, Any] = field(default_factory=dict)
    # Optional columns the workbook does not have: not provided, unlike a blank cell.
    absent: frozenset[str] = frozenset()

    def __getitem__(self, header: str) -> Any:
        return self.values.get(header)

    def provides(self, header: str) -> bool:
        return header not in self.absent


@dataclass
class ReadResult:
    rows: dict[Sheet, list[Row]]
    problems: list[Problem]
    absent_columns: dict[Sheet, list[str]]


class CellError(ValueError):
    pass


def read_workbook(path: Path) -> ReadResult:
    """Raises the openpyxl/zipfile errors for a file that is not an .xlsx workbook."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    result = ReadResult(rows={}, problems=[], absent_columns={})
    try:
        for sheet in DATA_SHEETS:
            if sheet.title not in workbook.sheetnames:
                result.problems.append(
                    Problem(sheet.title, "the sheet is missing from the workbook")
                )
                result.rows[sheet] = []
                continue
            _read_sheet(sheet, workbook[sheet.title].iter_rows(), result)
    finally:
        workbook.close()
    return result


def _read_sheet(sheet: Sheet, cells: Any, result: ReadResult) -> None:
    positions: dict[str, int] = {}
    absent: frozenset[str] = frozenset()
    rows = result.rows[sheet] = []
    for number, raw in enumerate(cells, start=HEADER_ROW):
        values = tuple(cell.value for cell in raw)
        if number == HEADER_ROW:
            found = _header_positions(sheet, values, result.problems)
            if found is None:
                return
            positions = found
            absent = frozenset(c.header for c in sheet.columns if c.header not in positions)
            if absent:
                result.absent_columns[sheet] = [
                    c.header for c in sheet.columns if c.header in absent
                ]
            continue
        if number < FIRST_DATA_ROW or all(_is_blank(value) for value in values):
            continue
        row = _read_row(sheet, number, values, positions, result.problems)
        row.absent = absent
        rows.append(row)


def _header_positions(
    sheet: Sheet, values: tuple[Any, ...], problems: list[Problem]
) -> dict[str, int] | None:
    """A missing optional column is left out of the positions: the rows do not provide
    it, so a workbook filled in before that column was added still imports without
    touching the field. Unknown columns are ignored."""
    found = {str(value).strip(): index for index, value in enumerate(values) if value}
    missing = [
        column.header
        for column in sheet.columns
        if (column.required or column.identifies) and column.header not in found
    ]
    for header in missing:
        problems.append(Problem(sheet.title, f'the column "{header}" is missing from row 1'))
    if missing:
        return None
    return {
        column.header: found[column.header] for column in sheet.columns if column.header in found
    }


def _read_row(
    sheet: Sheet,
    number: int,
    values: tuple[Any, ...],
    positions: dict[str, int],
    problems: list[Problem],
) -> Row:
    row = Row(sheet, number)
    for column in sheet.columns:
        position = positions.get(column.header)
        raw = values[position] if position is not None and position < len(values) else None
        if _is_blank(raw):
            if column.required:
                problems.append(Problem(sheet.title, "is required", number, column.header))
            continue
        try:
            row.values[column.header] = parse_cell(column, raw)
        except CellError as error:
            problems.append(Problem(sheet.title, str(error), number, column.header))
    return row


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def parse_cell(column: Column, raw: Any) -> Any:
    """Raises CellError with a message the client can act on."""
    match column.kind:
        case Kind.TEXT:
            text = _text(raw)
            if column.max_length and len(text) > column.max_length:
                raise CellError(f"is {len(text)} characters long; the limit is {column.max_length}")
            return text
        case Kind.WHOLE:
            return _whole_number(raw)
        case Kind.PRICE:
            return _price(raw)
        case Kind.YES_NO:
            return _yes_no(raw)
        case Kind.HEX:
            text = _text(raw)
            if not HEX_CODE.match(text):
                raise CellError(f'"{text}" is not a hex colour like #D8A47F')
            return text


def _text(raw: Any) -> str:
    # Excel stores a number typed into a text column, such as an SKU, as a float.
    if isinstance(raw, float) and raw.is_integer():
        return str(int(raw))
    return str(raw).strip()


def _whole_number(raw: Any) -> int:
    try:
        number = Decimal(_text(raw))
    except InvalidOperation:
        raise CellError(f'"{_text(raw)}" is not a whole number') from None
    if number != number.to_integral_value() or isinstance(raw, bool):
        raise CellError(f'"{_text(raw)}" is not a whole number')
    if number < 0:
        raise CellError(f"{number} is negative; use 0 or more")
    return int(number)


def _price(raw: Any) -> Decimal:
    text = _text(raw)
    if isinstance(raw, float):
        # repr() of a float is the shortest string that round-trips, so 1299.99 stays
        # 1299.99 rather than its binary expansion.
        text = repr(raw)
    try:
        price = Decimal(text)
    except InvalidOperation:
        raise CellError(
            f'"{text}" is not a price; write a plain number such as 2450, without commas or "Rs"'
        ) from None
    if not price.is_finite() or isinstance(raw, bool):
        raise CellError(f'"{text}" is not a price')
    if price < 0:
        raise CellError(f"{text} is negative")
    if price.as_tuple().exponent < -2:  # type: ignore[operator]  # finite, so an int
        raise CellError(f"{text} has more than 2 decimal places")
    if price > MAX_PRICE:
        raise CellError(f"{text} is larger than the shop can store")
    return price.quantize(Decimal("0.01"))


def _yes_no(raw: Any) -> bool:
    if isinstance(raw, bool):
        return raw
    text = _text(raw).lower()
    if text in TRUE_SPELLINGS:
        return True
    if text in FALSE_SPELLINGS:
        return False
    raise CellError(f'"{_text(raw)}" is not Yes or No')
